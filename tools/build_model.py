"""Genera research/05_MODELOS_DE_NEGOCIO.xlsx con fórmulas editables."""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUT = sys.argv[1]
wb = Workbook()

BOLD = Font(bold=True)
SUP = PatternFill("solid", fgColor="FFF2CC")   # supuesto
DAT = PatternFill("solid", fgColor="DDEBF7")   # dato con fuente
HDR = PatternFill("solid", fgColor="D9D9D9")

# ---------------------------------------------------------------- Léeme
ws = wb.active
ws.title = "Leeme"
lines = [
    "05 — Modelos de negocio (A5) · fecha 2026-10-19 · moneda USD",
    "",
    "Fórmula de ingresos visible: Ingresos por colocación = colocaciones (suscripciones) × ticket promedio × comisión total.",
    "Todas las celdas de 'Supuestos' son editables; 'Modelo_M2', 'Modelo_M1' y 'Sensibilidad' se recalculan con fórmulas.",
    "Celdas amarillas = SUPUESTO sin respaldo externo (NO es dato). Celdas celestes = DATO con fuente (ID de evidencia en 09_EVIDENCIAS.csv / frentes/A5A7_evidencia.csv).",
    "M2 = plataforma integral minorista (operada por/bajo entidad autorizada: PFC u operador con ALyC/fiduciario). Paga CAC de inversores.",
    "M1 = originador/marketplace B2B que deriva emisores a un operador autorizado; no capta inversores ni recibe fondos. La licitud de un fee de originación es HIPÓTESIS PARA ABOGADO.",
    "M0 (directorio informativo) y M3 (secundario) no se modelan numéricamente: M0 no tiene ingreso transaccional lícito identificado sin riesgo regulatorio; M3 no tiene volumen verificable (ver 10_DECISION.md).",
    "Simplificación conservadora: en M2 todos los inversores activos del año se tratan como nuevos (CAC completo). Sin impuestos (IIBB, IVA, Ganancias) ni inflación ni tipo de cambio.",
    "Conversión UVA→USD usada en fuentes: UVA ARS 2.159,10; USD mayorista ARS 1.514,57 (BCRA, E5043).",
]
for i, t in enumerate(lines, 1):
    ws.cell(row=i, column=1, value=t)
ws["A1"].font = Font(bold=True, size=13)
ws.column_dimensions["A"].width = 150

# ---------------------------------------------------------------- Supuestos
s = wb.create_sheet("Supuestos")
hdr = ["clave", "parámetro", "unidad", "conservador", "base", "optimista", "clase", "fuente / ancla"]
s.append(hdr)
for c in range(1, len(hdr) + 1):
    s.cell(row=1, column=c).font = BOLD
    s.cell(row=1, column=c).fill = HDR

params = [
    # clave, parámetro, unidad, cons, base, opt, clase, fuente
    ("campanas", "Campañas/emisiones cerradas por año", "n/año", 4, 12, 30, "SUPUESTO", "Sin ancla. Depende de originación (A1: solo 3 emisiones tokenizadas con oferta pública identificadas en AR)."),
    ("monto", "Monto medio colocado por campaña", "USD", 150000, 400000, 1000000, "SUPUESTO", "Tope PFC UVA 1,5M ≈ USD 2,14M por proyecto/12 meses (E5009, DATO)."),
    ("take_em", "Comisión de éxito al emisor (M2)", "% del monto", 0.04, 0.06, 0.08, "SUPUESTO", "Ancla internacional 6–7,9% (Wefunder E5001, Republic E5003)."),
    ("fee_inv", "Fee al inversor (M2)", "% del monto", 0.0, 0.01, 0.02, "SUPUESTO", "Ancla internacional 2–3,5% (E5002, E5005)."),
    ("ticket", "Ticket medio por suscripción", "USD", 300, 800, 2000, "SUPUESTO", "Tope RG 1125: UVA 3.000 por emisión (≈USD 4.300, E5011, sujeto del tope a confirmar)."),
    ("recurr", "Suscripciones por inversor activo por año", "n", 1, 2, 4, "SUPUESTO", "Sin ancla."),
    ("cac_inv", "CAC por inversor activo", "USD", 2000, 267, 50, "SUPUESTO", "Derivado de CAC registro 100/40/15 ÷ conversión 5/15/30% (A5 §2.6). Sin benchmark primario."),
    ("kyc", "Costo KYC por inversor", "USD", 1.9, 1.4, 0.8, "DATO", "Sumsub/Veriff precios públicos (E5018, E5019)."),
    ("atencion", "Atención por inversor activo/año", "USD", 12, 6, 3, "SUPUESTO", "Ancla débil Nubank (E5028)."),
    ("cac_em", "CAC por emisor que cierra", "USD", 10000, 4000, 1500, "SUPUESTO", "Sin fuente."),
    ("dd", "Costo propio de due diligence por emisor", "USD", 5000, 3000, 1500, "SUPUESTO", "Sin fuente pública; cotizar con estudios."),
    ("estruct", "Costo de estructuración pagado por el emisor (legal, fiduciario, CNV, auditoría)", "USD", 40000, 20000, 8000, "SUPUESTO", "Sin tarifa pública. Referencia: FF Total Assets V ≈2,18% s/VN sin impuestos (E5016), FF grande."),
    ("arancel_cnv", "Arancel CNV autorización FF pyme", "% del monto", 0.00015, 0.00015, 0.00015, "DATO", "Res. ME 1/2026 Anexo II: 0,015% FF pyme (E5012)."),
    ("prov", "Provisión por reclamos/fraude/default (costo plataforma)", "% del GMV", 0.005, 0.0025, 0.001, "SUPUESTO", "Sin fuente. Riesgo reputacional ver 08."),
    ("fijo_m2", "Costo fijo anual M2 (equipo, compliance, tecnología, tasas)", "USD/año", 600000, 350000, 200000, "SUPUESTO", "Incluye tasa CNV si aplica (PSAV PJ ARS 13,36M ≈ USD 8.821, E5013)."),
    ("fee_orig", "Fee de originación cobrado al operador/emisor (M1)", "% del monto", 0.015, 0.025, 0.035, "SUPUESTO", "Sin ancla local. Licitud: HIPÓTESIS PARA ABOGADO (ver 11)."),
    ("fee_prep", "Fee fijo de preparación de legajo por emisor (M1)", "USD", 2000, 4000, 6000, "SUPUESTO", "Sin ancla. Ref. briq.mx cobra dictamen legal MXN 17.000 (A2)."),
    ("fijo_m1", "Costo fijo anual M1 (equipo chico, sin captación)", "USD/año", 250000, 150000, 90000, "SUPUESTO", "Sin fuente."),
]
ROW = {}
for i, p in enumerate(params, start=2):
    s.append(list(p))
    ROW[p[0]] = i
    fill = SUP if p[6] == "SUPUESTO" else DAT
    for c in (4, 5, 6):
        s.cell(row=i, column=c).fill = fill
        if "%" in p[2]:
            s.cell(row=i, column=c).number_format = "0.00%"
        else:
            s.cell(row=i, column=c).number_format = "#,##0.00"
for col, w in zip("ABCDEFGH", (12, 62, 14, 14, 14, 14, 11, 95)):
    s.column_dimensions[col].width = w

SCEN = {"conservador": "D", "base": "E", "optimista": "F"}


def ref(key, col):
    return f"Supuestos!${col}${ROW[key]}"


def build_model(name, rows):
    m = wb.create_sheet(name)
    m.append(["concepto", "conservador", "base", "optimista", "fórmula (texto)"])
    for c in range(1, 6):
        m.cell(row=1, column=c).font = BOLD
        m.cell(row=1, column=c).fill = HDR
    pos = {}
    for r, (key, label, fmt, expr, txt) in enumerate(rows, start=2):
        pos[key] = r
        m.cell(row=r, column=1, value=label)
        m.cell(row=r, column=5, value=txt)
        for j, (sc, scol) in enumerate(SCEN.items()):
            col = get_column_letter(2 + j)
            f = expr(lambda k: ref(k, scol), lambda k: f"{col}{pos[k]}")
            cell = m.cell(row=r, column=2 + j, value=f)
            cell.number_format = fmt
        if key.startswith("res") or key.startswith("ing_tot"):
            m.cell(row=r, column=1).font = BOLD
    m.column_dimensions["A"].width = 58
    for c in "BCD":
        m.column_dimensions[c].width = 16
    m.column_dimensions["E"].width = 90
    return pos


USD = "#,##0"
PCT = "0.0%"
NUM = "#,##0.0"

m2_rows = [
    ("gmv", "GMV (monto colocado)", USD, lambda S, M: f"={S('campanas')}*{S('monto')}", "campañas × monto medio"),
    ("susc", "Colocaciones (suscripciones)", USD, lambda S, M: f"={M('gmv')}/{S('ticket')}", "GMV ÷ ticket"),
    ("tk", "Ticket promedio", USD, lambda S, M: f"={S('ticket')}", "supuesto"),
    ("com", "Comisión total (emisor + inversor)", "0.00%", lambda S, M: f"={S('take_em')}+{S('fee_inv')}", "take emisor + fee inversor"),
    ("ing_tot", "INGRESOS = colocaciones × ticket × comisión", USD, lambda S, M: f"={M('susc')}*{M('tk')}*{M('com')}", "fórmula del mandato"),
    ("inv_act", "Inversores activos", USD, lambda S, M: f"={M('susc')}/{S('recurr')}", "suscripciones ÷ recurrencia"),
    ("c_cacinv", "Costo adquisición inversores", USD, lambda S, M: f"={M('inv_act')}*{S('cac_inv')}", "inversores × CAC (todos tratados como nuevos)"),
    ("c_kyc", "Costo KYC", USD, lambda S, M: f"={M('inv_act')}*{S('kyc')}", "inversores × KYC"),
    ("c_at", "Costo atención", USD, lambda S, M: f"={M('inv_act')}*{S('atencion')}", "inversores × atención"),
    ("c_cacem", "Costo adquisición emisores", USD, lambda S, M: f"={S('campanas')}*{S('cac_em')}", "campañas × CAC emisor"),
    ("c_dd", "Costo due diligence", USD, lambda S, M: f"={S('campanas')}*{S('dd')}", "campañas × DD"),
    ("c_prov", "Provisión reclamos/fraude", USD, lambda S, M: f"={M('gmv')}*{S('prov')}", "GMV × provisión"),
    ("c_fijo", "Costo fijo anual", USD, lambda S, M: f"={S('fijo_m2')}", "supuesto"),
    ("c_tot", "COSTOS TOTALES", USD, lambda S, M: f"=SUM({M('c_cacinv')}:{M('c_fijo')})", "suma"),
    ("res", "RESULTADO OPERATIVO (antes de impuestos)", USD, lambda S, M: f"={M('ing_tot')}-{M('c_tot')}", "ingresos − costos"),
    ("res_m", "Margen operativo", PCT, lambda S, M: f"=IF({M('ing_tot')}=0,0,{M('res')}/{M('ing_tot')})", "resultado ÷ ingresos"),
    ("cv", "Costo variable por USD de GMV", "0.000", lambda S, M: f"=({S('cac_inv')}+{S('kyc')}+{S('atencion')})/({S('ticket')}*{S('recurr')})+({S('cac_em')}+{S('dd')})/{S('monto')}+{S('prov')}", "(CAC+KYC+atención)/(ticket×recurrencia) + (CAC em.+DD)/monto + provisión"),
    ("mc", "Margen de contribución por USD de GMV", "0.000", lambda S, M: f"={M('com')}-{M('cv')}", "comisión − costo variable"),
    ("be", "GMV de equilibrio (si margen de contribución > 0)", USD, lambda S, M: f"=IF({M('mc')}<=0,\"NUNCA (contribución negativa)\",{S('fijo_m2')}/{M('mc')})", "fijo ÷ margen de contribución"),
    ("be_c", "Campañas de equilibrio", NUM, lambda S, M: f"=IF(ISNUMBER({M('be')}),{M('be')}/{S('monto')},\"n/a\")", "GMV equilibrio ÷ monto"),
    ("costo_em", "Costo total para el emisor (% del monto)", PCT, lambda S, M: f"={S('take_em')}+{S('estruct')}/{S('monto')}+{S('arancel_cnv')}", "take + estructuración/monto + arancel CNV"),
]
build_model("Modelo_M2", m2_rows)

m1_rows = [
    ("gmv", "GMV originado y colocado por el operador", USD, lambda S, M: f"={S('campanas')}*{S('monto')}", "campañas × monto medio"),
    ("susc", "Colocaciones (emisiones originadas)", NUM, lambda S, M: f"={S('campanas')}", "en M1 la 'colocación' es la emisión derivada"),
    ("tk", "Ticket promedio (monto por emisión)", USD, lambda S, M: f"={S('monto')}", "supuesto"),
    ("com", "Comisión de originación", "0.00%", lambda S, M: f"={S('fee_orig')}", "supuesto — licitud a confirmar"),
    ("ing_var", "Ingresos variables = colocaciones × ticket × comisión", USD, lambda S, M: f"={M('susc')}*{M('tk')}*{M('com')}", "fórmula del mandato"),
    ("ing_prep", "Ingresos por preparación de legajo", USD, lambda S, M: f"={S('campanas')}*{S('fee_prep')}*1", "campañas × fee fijo (supone que solo cobran las que cierran: conservador)"),
    ("ing_tot", "INGRESOS TOTALES", USD, lambda S, M: f"={M('ing_var')}+{M('ing_prep')}", "suma"),
    ("c_cacem", "Costo adquisición emisores", USD, lambda S, M: f"={S('campanas')}*{S('cac_em')}", "campañas × CAC emisor"),
    ("c_dd", "Costo due diligence", USD, lambda S, M: f"={S('campanas')}*{S('dd')}", "campañas × DD"),
    ("c_prov", "Provisión reclamos (reputacional)", USD, lambda S, M: f"={M('gmv')}*{S('prov')}/2", "GMV × provisión ÷ 2 (no capta fondos; supuesto)"),
    ("c_fijo", "Costo fijo anual", USD, lambda S, M: f"={S('fijo_m1')}", "supuesto"),
    ("c_tot", "COSTOS TOTALES", USD, lambda S, M: f"=SUM({M('c_cacem')}:{M('c_fijo')})", "suma"),
    ("res", "RESULTADO OPERATIVO (antes de impuestos)", USD, lambda S, M: f"={M('ing_tot')}-{M('c_tot')}", "ingresos − costos"),
    ("res_m", "Margen operativo", PCT, lambda S, M: f"=IF({M('ing_tot')}=0,0,{M('res')}/{M('ing_tot')})", "resultado ÷ ingresos"),
    ("mc_c", "Contribución por campaña", USD, lambda S, M: f"={S('monto')}*{S('fee_orig')}+{S('fee_prep')}-{S('cac_em')}-{S('dd')}-{S('monto')}*{S('prov')}/2", "monto×fee + prep − CAC − DD − provisión"),
    ("be_c", "Campañas de equilibrio", NUM, lambda S, M: f"=IF({M('mc_c')}<=0,\"NUNCA\",{S('fijo_m1')}/{M('mc_c')})", "fijo ÷ contribución por campaña"),
    ("costo_em", "Costo total para el emisor (% del monto)", PCT, lambda S, M: f"=({S('fee_orig')}*{S('monto')}+{S('fee_prep')}+{S('estruct')})/{S('monto')}+{S('arancel_cnv')}+0.02", "fee orig. + prep + estructuración + arancel + 2% colocación del operador (SUPUESTO)"),
]
build_model("Modelo_M1", m1_rows)

# ---------------------------------------------------------------- Sensibilidad
z = wb.create_sheet("Sensibilidad")
B = lambda k: ref(k, "E")  # escenario base
z["A1"] = "Sensibilidad M2 (resto de supuestos en escenario BASE). Resultado operativo anual en USD."
z["A1"].font = BOLD
z["A3"] = "Tabla 1: GMV anual (filas) × comisión total (columnas)"
z["A3"].font = BOLD
takes = [0.03, 0.04, 0.05, 0.06, 0.07, 0.08, 0.10]
gmvs = [1e6, 2.5e6, 5e6, 10e6, 20e6, 40e6]
z["A4"] = "GMV \\ comisión"
for j, t in enumerate(takes):
    c = z.cell(row=4, column=2 + j, value=t)
    c.number_format = "0%"
    c.font = BOLD
cv_expr = f"(({B('cac_inv')}+{B('kyc')}+{B('atencion')})/({B('ticket')}*{B('recurr')})+({B('cac_em')}+{B('dd')})/{B('monto')}+{B('prov')})"
for i, g in enumerate(gmvs):
    r = 5 + i
    z.cell(row=r, column=1, value=g).number_format = USD
    for j in range(len(takes)):
        col = get_column_letter(2 + j)
        z.cell(row=r, column=2 + j, value=f"=$A{r}*({col}$4-{cv_expr})-{B('fijo_m2')}").number_format = USD

z["A13"] = "Tabla 2: ticket medio (filas) × suscripciones por inversor/año (columnas) → resultado con GMV y comisión base"
z["A13"].font = BOLD
tickets = [300, 500, 800, 1500, 3000, 4300]
recs = [1, 2, 3, 5, 8]
z["A14"] = "ticket \\ recurrencia"
for j, rc in enumerate(recs):
    z.cell(row=14, column=2 + j, value=rc).font = BOLD
gmv_b = f"({B('campanas')}*{B('monto')})"
com_b = f"({B('take_em')}+{B('fee_inv')})"
for i, tk in enumerate(tickets):
    r = 15 + i
    z.cell(row=r, column=1, value=tk).number_format = USD
    for j in range(len(recs)):
        col = get_column_letter(2 + j)
        cv = f"(({B('cac_inv')}+{B('kyc')}+{B('atencion')})/($A{r}*{col}$14)+({B('cac_em')}+{B('dd')})/{B('monto')}+{B('prov')})"
        z.cell(row=r, column=2 + j, value=f"={gmv_b}*({com_b}-{cv})-{B('fijo_m2')}").number_format = USD

z["A23"] = "Tabla 3: CAC por inversor activo (filas) × ticket (columnas) → margen de contribución por USD de GMV (recurrencia base)"
z["A23"].font = BOLD
cacs = [20, 50, 100, 267, 500, 1000]
z["A24"] = "CAC \\ ticket"
for j, tk in enumerate(tickets):
    z.cell(row=24, column=2 + j, value=tk).font = BOLD
for i, cac in enumerate(cacs):
    r = 25 + i
    z.cell(row=r, column=1, value=cac).number_format = USD
    for j in range(len(tickets)):
        col = get_column_letter(2 + j)
        z.cell(row=r, column=2 + j,
               value=f"={com_b}-(($A{r}+{B('kyc')}+{B('atencion')})/({col}$24*{B('recurr')})+({B('cac_em')}+{B('dd')})/{B('monto')}+{B('prov')})").number_format = "0.0%"

z["A33"] = "Tabla 4: M1 — campañas/año (filas) × monto medio (columnas) → resultado operativo M1 (fees base)"
z["A33"].font = BOLD
camps = [2, 4, 8, 12, 20, 30]
montos = [100000, 250000, 500000, 1000000, 2000000]
z["A34"] = "campañas \\ monto"
for j, mt in enumerate(montos):
    c = z.cell(row=34, column=2 + j, value=mt)
    c.number_format = USD
    c.font = BOLD
for i, cp in enumerate(camps):
    r = 35 + i
    z.cell(row=r, column=1, value=cp)
    for j in range(len(montos)):
        col = get_column_letter(2 + j)
        z.cell(row=r, column=2 + j,
               value=f"=$A{r}*({col}$34*{B('fee_orig')}+{B('fee_prep')}-{B('cac_em')}-{B('dd')}-{col}$34*{B('prov')}/2)-{B('fijo_m1')}").number_format = USD
z.column_dimensions["A"].width = 22
for j in range(2, 10):
    z.column_dimensions[get_column_letter(j)].width = 15

wb.save(OUT)
print("ok", OUT)
